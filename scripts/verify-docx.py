#!/usr/bin/env python3
"""Verify the generated DOCX structure."""

from docx import Document

doc = Document("laporan_magang_final.docx")

paras = len(doc.paragraphs)
tables = len(doc.tables)
sections = len(doc.sections)

# Count inline images
img_count = 0
ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
drawing_ns = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"

for p in doc.paragraphs:
    for r in p.runs:
        drawings = r._element.findall(f".//{drawing_ns}inline")
        img_count += len(drawings)

print(f"Sections: {sections}")
print(f"Paragraphs: {paras}")
print(f"Tables: {tables}")
print(f"Images found (inline elements): {img_count}")

print("\n--- First 15 paragraphs ---")
for i, p in enumerate(doc.paragraphs[:15]):
    txt = p.text[:80] if p.text else "(empty)"
    style = p.style.name if p.style else "None"
    print(f"  [{i:3d}] ({style:25s}) {txt}")

print("\n--- Heading paragraphs ---")
heading_count = 0
for i, p in enumerate(doc.paragraphs):
    if p.style and "Heading" in p.style.name:
        heading_count += 1
        if heading_count <= 25:
            txt = p.text[:70] if p.text else "(empty)"
            print(f"  [{i:3d}] ({p.style.name:15s}) {txt}")
print(f"\nTotal headings: {heading_count}")

print("\n--- Last 10 paragraphs ---")
for i, p in enumerate(doc.paragraphs[-10:]):
    txt = p.text[:80] if p.text else "(empty)"
    print(f"  [{paras-10+i:3d}] {txt}")

print("\n--- Tables ---")
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first = t.rows[0].cells[0].text[:50] if rows > 0 and cols > 0 else "(empty)"
    print(f"  Table {i}: {rows}r x {cols}c -> \"{first}\"")
